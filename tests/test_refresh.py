"""Yenileme hataları çalışan veriyi bozmamalı; tarih gerçek kaynak kontrolünü anlatmalı."""
import copy
import hashlib
import importlib.util
import io
import json
import subprocess
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
spec = importlib.util.spec_from_file_location("refresh_build", ROOT / "scripts/build_data.py")
bd = importlib.util.module_from_spec(spec)
spec.loader.exec_module(bd)


class RefreshTransaction(unittest.TestCase):
    """İkinci kurum kırılırsa ilk kurumun dosyaları da eski hâlinde kalır."""
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.site = self.root / "site"
        self.site.mkdir()
        self.record = json.loads((ROOT / "site/data-maku.json").read_text())["agreements"][0]
        wb = openpyxl.Workbook()
        wb.active.append(["fixture"])
        stream = io.BytesIO()
        wb.save(stream)
        self.data = stream.getvalue()
        self.digest = hashlib.sha256(self.data).hexdigest()
        self.unis, sources, registry = [], [], []
        for uid in ("alpha", "beta"):
            local = self.root / "lokal" / uid / "source.xlsx"
            local.parent.mkdir(parents=True)
            local.write_bytes(self.data)
            uni = dict(id=uid, nameTr=uid, nameEn=uid, abbr=uid, monogram=uid,
                       creditTr=uid, creditEn=uid, hasGuide=False, sheet=0,
                       listUrl="https://example.edu/agreements", local=local,
                       parser=lambda ws: [copy.deepcopy(self.record)], downloadMode="page-xlsx")
            self.unis.append(uni)
            sources.append(dict(id=uid, sha256=self.digest, indirme_tarihi="2026-01-01",
                                son_kontrol_tarihi="2026-01-02", yerel_ad=str(local.relative_to(self.root))))
            registry.append(dict(id=uid, generatedAt="2026-01-03", hasGuide=False))
            (self.site / f"data-{uid}.json").write_text(json.dumps({"agreements": [self.record], "generatedAt": "2026-01-03"}))
        (self.site / "universities.json").write_text(json.dumps(registry))
        (self.site / "kaynak-kunyesi.json").write_text(json.dumps({"kaynaklar": sources}))
        for name, value in (("ROOT", self.root), ("SITE", self.site), ("UNIVERSITIES", self.unis)):
            p = patch.object(bd, name, value)
            p.start()
            self.addCleanup(p.stop)
        self.download = patch.object(bd, "download_source", return_value=(self.data, "https://example.edu/source.xlsx", None)).start()
        self.addCleanup(patch.stopall)

    def files(self):
        return {str(p.relative_to(self.root)): p.read_bytes() for p in self.root.rglob("*") if p.is_file()}

    def test_dry_run_preserves_every_source_and_output(self):
        before = self.files()
        bd.build(pull=True, dry_run=True)
        self.assertEqual(before, self.files())

    def test_second_download_failure_preserves_first_university(self):
        before = self.files()
        self.download.side_effect = [(self.data, "https://example.edu/source.xlsx", None), OSError("offline")]
        with self.assertRaises(OSError):
            bd.build(pull=True)
        self.assertEqual(before, self.files())

    def test_empty_parser_preserves_existing_outputs(self):
        before = self.files()
        self.unis[1]["parser"] = lambda ws: []
        with self.assertRaisesRegex(ValueError, "hiç kayıt"):
            bd.build(pull=True)
        self.assertEqual(before, self.files())

    def test_selected_university_preserves_other_data_and_timestamp(self):
        before = (self.site / "data-beta.json").read_bytes()
        bd.build(pull=True, university_ids=["alpha"])
        self.assertEqual(before, (self.site / "data-beta.json").read_bytes())
        registry = json.loads((self.site / "universities.json").read_text())
        self.assertEqual(registry[1]["generatedAt"], "2026-01-03")
        self.assertEqual(registry[0]["sourceCheckedAt"], date.today().isoformat())
        self.assertEqual(registry[0]["sourceDownloadedAt"], "2026-01-01")

    def test_offline_build_does_not_claim_fresh_check(self):
        bd.build(pull=False, university_ids=["alpha"])
        data = json.loads((self.site / "data-alpha.json").read_text())
        self.assertEqual(data["sourceCheckedAt"], "2026-01-02")
        self.assertEqual(data["source"], "local-snapshot")
        self.download.assert_not_called()

    def test_modified_local_file_needs_provenance(self):
        self.unis[0]["local"].write_bytes(b"changed")
        before = self.files()
        with self.assertRaisesRegex(ValueError, "uyuşmuyor"):
            bd.build()
        self.assertEqual(before, self.files())

    def test_unknown_university_does_not_start_download(self):
        with self.assertRaisesRegex(ValueError, "Bilinmeyen kurum"):
            bd.build(pull=True, university_ids=["typo"])
        self.download.assert_not_called()


class DownloadSelection(unittest.TestCase):
    """Birden fazla kaynak varsa hangisinin güncel olduğu tahmin edilmez."""
    def test_relative_escaped_links_are_resolved_and_deduplicated(self):
        links = bd.SourceLinks("https://example.edu/list")
        links.feed('<a href="/a.xlsx?x=1&amp;y=2">a</a><a href="/a.xlsx?x=1&amp;y=2">b</a>')
        self.assertEqual(links.links, {"https://example.edu/a.xlsx?x=1&y=2"})

    def test_multiple_workbook_links_stop_download(self):
        with patch.object(bd, "fetch_bytes", return_value=(b'<a href="a.xlsx">a</a><a href="b.xlsx">b</a>', "https://example.edu/list")):
            with self.assertRaisesRegex(ValueError, "2 Excel"):
                bd.download_source(dict(id="test", downloadMode="page-xlsx", listUrl="https://example.edu/list"))

    def test_archive_requires_exactly_one_workbook(self):
        stream = io.BytesIO()
        with ZipFile(stream, "w") as archive:
            archive.writestr("one.xlsx", b"one")
            archive.writestr("two.xlsx", b"two")
        with patch.object(bd, "fetch_bytes", return_value=(stream.getvalue(), "https://example.edu/download")):
            with self.assertRaisesRegex(ValueError, "2 Excel"):
                bd.download_source(dict(id="test", downloadMode="share-zip", pullUrl="https://example.edu/download"))

    def test_html_response_is_not_accepted_as_workbook(self):
        with patch.object(bd, "fetch_bytes", return_value=(b"<html>Login</html>", "https://example.edu/login")):
            with self.assertRaisesRegex(ValueError, "Excel dosyası değil"):
                bd.download_source(dict(id="test", downloadMode="direct", pullUrl="https://example.edu/file"))


class BilecikMergedSource(unittest.TestCase):
    """Birleşik kontenjan ve boş bölüm satırları yanlış kesinlik üretmemeli."""
    def sheet(self):
        ws = openpyxl.Workbook().active
        for col, label in bd.BILECIK_HEADERS.items():
            ws.cell(2, col, label)
        for col, val in {2: "2026/2029", 6: "Test University", 7: "Polonya/Test", 8: "PL TEST01", 9: "061-ICT", 10: "Mühendislik", 11: "Bilgisayar", 12: "YL", 14: "2*5"}.items():
            ws.cell(3, col, val)
        ws.cell(4, 9, "071-Engineering")
        ws.cell(4, 11, "Elektrik")
        for col in (2, 6, 7, 8, 10, 12, 14):
            ws.merge_cells(start_row=3, end_row=4, start_column=col, end_column=col)
        return ws

    def test_shared_quota_keeps_duration_and_marks_both_rows(self):
        rows = bd.parse_bilecik(self.sheet())
        self.assertEqual(len(rows), 2)
        self.assertEqual([r["department"] for r in rows], ["Bilgisayar", "Elektrik"])
        self.assertTrue(all(r["sharedQuota"] and r["quotaStudy"] == "2*5" for r in rows))

    def test_blank_department_does_not_inherit_previous_department(self):
        ws = self.sheet()
        ws.cell(4, 11).value = None
        row = bd.parse_bilecik(ws)[1]
        self.assertEqual(row["department"], "Mühendislik")
        self.assertEqual(row["sourceDiff"]["department"]["sebep"], "fakulte-adi")

    def test_multiple_families_are_not_collapsed_into_one(self):
        ws = self.sheet()
        ws.cell(3, 9).value = "0232-Literature 0114-Teaching"
        row = bd.parse_bilecik(ws)[0]
        self.assertIsNone(row["iscedFamily"])
        self.assertIn("0114", row["sourceField"])

    def test_nonstandard_code_is_preserved_without_guessing_family(self):
        ws = self.sheet()
        ws.cell(3, 9).value = "688-Informatyka"
        row = bd.parse_bilecik(ws)[0]
        self.assertIsNone(row["iscedFamily"])
        self.assertEqual(row["sourceField"], "688-Informatyka")

    def test_header_shift_stops_parser(self):
        ws = self.sheet()
        ws.cell(2, 14).value = "New column"
        with self.assertRaisesRegex(ValueError, "başlığı değişti"):
            bd.parse_bilecik(ws)

    def test_contact_column_is_never_read(self):
        """İletişim sütunu yalnız çıktıda gizlenmemeli; ayrıştırıcı onu okumamalı."""
        ws = self.sheet()
        original_cell = ws.cell

        def guarded_cell(row, column, *args, **kwargs):
            self.assertNotEqual(column, 18, "Bilecik iletişim sütunu okunmamalı")
            return original_cell(row, column, *args, **kwargs)

        with patch.object(ws, "cell", side_effect=guarded_cell):
            rows = bd.parse_bilecik(ws)
        self.assertEqual(len(rows), 2)

    def test_single_level_unmerged_quota_is_not_marked_shared(self):
        ws = self.sheet()
        ws.unmerge_cells("N3:N4")
        ws.unmerge_cells("L3:L4")
        ws.cell(3, 12).value = 0
        self.assertEqual(bd.parse_bilecik(ws)[0]["levels"], {"lisans": "2"})


class PublishedProvenance(unittest.TestCase):
    def test_registry_data_and_catalogue_agree(self):
        registry = json.loads((ROOT / "site/universities.json").read_text())
        catalogue = {s["id"]: s for s in json.loads((ROOT / "site/kaynak-kunyesi.json").read_text())["kaynaklar"]}
        for uni in registry:
            data = json.loads((ROOT / "site" / f"data-{uni['id']}.json").read_text())
            source = catalogue[uni["id"]]
            for field, origin in (("sourceCheckedAt", "son_kontrol_tarihi"), ("sourceDownloadedAt", "indirme_tarihi"), ("sourceSha256", "sha256")):
                self.assertEqual(uni[field], source[origin])
                self.assertEqual(data[field], source[origin])
            self.assertEqual(data["count"], source["kayit_sayisi"])

    def test_source_dates_and_stale_warning_in_both_languages(self):
        script = '''
const fs = require('fs'), vm = require('vm'), assert = require('assert');
const context = {window: {}, Date, Intl}; vm.createContext(context);
vm.runInContext(fs.readFileSync('site/veri.js', 'utf8'), context);
const info = context.window.atlasSourceInfo, now = Date.UTC(2026, 8, 20);
for (const lang of ['tr', 'en']) {
  const old = info({sourceCheckedAt: '2026-07-07', generatedAt: '2026-09-20'}, lang, now);
  assert(old.stale); assert(old.statusText); assert.match(old.dateText, /7 (Temmuz|July) 2026/);
  const fresh = info({sourceCheckedAt: '2026-09-20', sourceDownloadedAt: '2026-07-07'}, lang, now);
  assert(!fresh.stale); assert.strictEqual(fresh.statusText, '');
  assert(info({generatedAt: '2026-09-20'}, lang, now).stale);
  assert(info({sourceCheckedAt: '2026-02-31'}, lang, now).stale);
  assert(info({sourceCheckedAt: '2027-01-01'}, lang, now).stale);
  assert(info({sourceDownloadedAt: '2026-07-07'}, lang, now).stale);
  assert(!info({sourceCheckedAt: '2026-09-20'}, lang, Date.UTC(2026, 8, 19, 23)).stale);
}
'''
        # A missing Node installation must fail, rather than silently skip UI logic.
        result = subprocess.run(["node", "-e", script], cwd=ROOT, capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)


class StaffQuotaCards(unittest.TestCase):
    """MAKÜ'nün iki personel sütunu ve ortak kontenjan açıklaması kartta kaybolmamalı."""

    def test_teaching_and_training_are_separate_and_keep_source_notes(self):
        script = '''
const fs = require('fs'), vm = require('vm'), assert = require('assert');
const context = {
  URLSearchParams, location: {search: '?uni=maku'},
  localStorage: {getItem: () => null}, atlasJson: () => new Promise(() => {})
};
vm.createContext(context);
vm.runInContext(fs.readFileSync('site/app.js', 'utf8'), context);
const rows = JSON.parse(fs.readFileSync('site/data-maku.json', 'utf8')).agreements;
context.record = rows.find(a => a.erasmusCode === 'BG SOFIA22' && a.department === 'Bilgisayar Mühendisliği');
assert(context.record, 'Gerçek MAKÜ örneği bulunamadı; kaynak değişmiş olabilir.');
assert.strictEqual(context.record.quotaStaffTeach, '2');
assert.strictEqual(context.record.quotaStaffTrain, '5 (anlaşmaya dalil olan 4 bölüm için toplam kontenjan)');
for (const [language, teaching, training] of [['tr', 'Ders verme', 'Eğitim alma'], ['en', 'Teaching', 'Training']]) {
  vm.runInContext(`lang = '${language}'`, context);
  const html = vm.runInContext('card(record)', context);
  assert(html.includes(teaching + ' · 2'), 'Ders verme kontenjanı ayrı görünmeli.');
  assert(html.includes(training + ' · ' + context.record.quotaStaffTrain), 'Eğitim alma kontenjanı ve ortaklık notu korunmalı.');
  const original = context.record;
  context.record = {...original, quotaStaffTeach: '0'};
  const trainingOnly = vm.runInContext('card(record)', context);
  assert(!trainingOnly.includes(teaching + ' · 0'));
  assert(trainingOnly.includes(training + ' · ' + original.quotaStaffTrain));
  context.record = original;
}
'''
        result = subprocess.run(["node", "-e", script], cwd=ROOT, capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)


if __name__ == "__main__":
    unittest.main()
